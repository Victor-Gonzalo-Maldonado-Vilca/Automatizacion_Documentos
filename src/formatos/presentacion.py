"""Formato: presentación académica (PPTX y PDF) con temas de color.

Es independiente de los demás formatos: la portada usa los datos del propio
presentacion.yaml. Opcionalmente puede tomarlos de un informe con `desde:`.
"""

from __future__ import annotations

import math
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from bib import formato_ieee, leer_bib
from resaltado import resaltar
from datos import CONFIG_INSTITUCION, RAIZ, ErrorDatos, a_lista, combinar, leer_yaml, slug

NOMBRE = "presentacion"
DESCRIPCION = "Presentación académica con temas de color"
ARCHIVO = "presentacion.yaml"
PLANTILLA_NUEVO = RAIZ / "plantillas" / "presentacion" / "nuevo"
SALIDAS = ("pptx", "pdf")
TEMAS = RAIZ / "plantillas" / "presentacion" / "temas.yaml"
ANCHO, ALTO = 13.333, 7.5          # 16:9 en pulgadas
MX = 0.6                           # margen lateral
TOPE_CUERPO, FONDO_CUERPO = 1.75, 6.75


# ============================================================================
#  Utilidades de bajo nivel
# ============================================================================

def _rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa.lstrip("#").upper())


def _mezclar(a: str, b: str, t: float) -> str:
    """Interpola dos colores hex (t=0 -> a, t=1 -> b)."""
    ca, cb = _rgb(a), _rgb(b)
    return "#" + "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ca, cb))


_MENORES = {"de", "del", "la", "las", "el", "los", "y", "e", "en", "a", "para", "por", "con"}


def _titulo_es(texto: str) -> str:
    """'UNIVERSIDAD NACIONAL DE SAN AGUSTIN' -> 'Universidad Nacional de San Agustin'."""
    palabras = texto.lower().split()
    return " ".join(w if (k and w in _MENORES) else w.capitalize() for k, w in enumerate(palabras))


def _num(valor, defecto: float) -> float:
    try:
        return float(valor)
    except (TypeError, ValueError):
        return defecto


def _verdad(valor, defecto: bool = True) -> bool:
    if valor in (None, ""):
        return defecto
    return str(valor).strip().lower() in ("true", "si", "sí", "yes", "1")


class Lienzo:
    """Envuelve una diapositiva con ayudas de dibujo en pulgadas."""

    def __init__(self, slide, tema: dict):
        self.s = slide
        self.t = tema

    def rect(self, x, y, w, h, color, forma=MSO_SHAPE.RECTANGLE, borde: str | None = None):
        sh = self.s.shapes.add_shape(forma, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.shadow.inherit = False
        if color:
            sh.fill.solid()
            sh.fill.fore_color.rgb = _rgb(color)
        else:
            sh.fill.background()
        if borde:
            sh.line.color.rgb = _rgb(borde)
            sh.line.width = Pt(0.75)
        else:
            sh.line.fill.background()
        return sh

    def texto(self, x, y, w, h, contenido="", tam=18, color=None, negrita=False, cursiva=False,
              fuente=None, alinear=PP_ALIGN.LEFT, ancla=MSO_ANCHOR.TOP, espaciado=None,
              interlinea=1.1):
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = ancla
        p = tf.paragraphs[0]
        p.alignment = alinear
        p.line_spacing = interlinea
        if contenido:
            self.run(p, contenido, tam, color, negrita, cursiva, fuente, espaciado)
        return tf

    def run(self, p, contenido, tam=18, color=None, negrita=False, cursiva=False, fuente=None, espaciado=None):
        r = p.add_run()
        r.text = str(contenido)
        f = r.font
        f.size = Pt(tam)
        f.bold = negrita
        f.italic = cursiva
        f.name = fuente or self.t["fuente_texto"]
        f.color.rgb = _rgb(color or self.t["texto"])
        if espaciado is not None:  # espaciado entre letras (en centésimas de pt)
            r._r.get_or_add_rPr().set("spc", str(int(espaciado * 100)))
        return r

    def imagen(self, ruta: Path, x, y, w, h, borde: str | None = None):
        """Inserta una imagen ajustada (sin deformar) y centrada en la caja."""
        if not ruta.exists():
            raise ErrorDatos(f"No existe la imagen: {ruta}")
        from PIL import Image
        with Image.open(ruta) as im:
            iw, ih = im.size
        escala = min(w / iw, h / ih)
        ww, hh = iw * escala, ih * escala
        px, py = x + (w - ww) / 2, y + (h - hh) / 2
        pic = self.s.shapes.add_picture(str(ruta), Inches(px), Inches(py), Inches(ww), Inches(hh))
        if borde:
            pic.line.color.rgb = _rgb(borde)
            pic.line.width = Pt(0.75)
        return pic, (px, py, ww, hh)


def _vineta(p, caracter: str, color: str, margen: float, sangria: float, escala: int = 100):
    """Viñeta personalizada (python-pptx no la expone directamente)."""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(margen)))
    pPr.set("indent", str(-Inches(sangria)))
    for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buNone", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    clr = etree.SubElement(pPr, qn("a:buClr"))
    etree.SubElement(clr, qn("a:srgbClr")).set("val", color.lstrip("#").upper())
    etree.SubElement(pPr, qn("a:buSzPct")).set("val", str(escala * 1000))
    etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
    etree.SubElement(pPr, qn("a:buChar")).set("char", caracter)


def _sin_vineta(p):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", "0")
    pPr.set("indent", "0")
    etree.SubElement(pPr, qn("a:buNone"))


def _lineas_estimadas(texto: str, ancho_in: float, tam: float) -> int:
    """Líneas aproximadas que ocupa un texto (Calibri ~0.5 em por carácter)."""
    por_linea = max(8, int(ancho_in * 72 / (tam * 0.5)))
    return max(1, math.ceil(len(texto) / por_linea))


def _normalizar_puntos(puntos) -> list[tuple[int, str]]:
    """Convierte la lista YAML en [(nivel, texto)]. Admite:
        - Punto simple
        - texto: Punto con subpuntos
          sub: [a, b]
    """
    salida = []
    for p in a_lista(puntos):
        if isinstance(p, dict):
            salida.append((0, p.get("texto", "")))
            for s in a_lista(p.get("sub")):
                salida.append((1, s if isinstance(s, str) else s.get("texto", "")))
        else:
            salida.append((0, p))
    return salida


def _escribir_markup(lz: Lienzo, p, texto: str, tam: float, color: str, fuente=None):
    """Soporta **negrita** y `código` dentro de los puntos."""
    import re
    partes = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto)
    for parte in partes:
        if not parte:
            continue
        if parte.startswith("**"):
            lz.run(p, parte[2:-2], tam, lz.t["texto"] if color == lz.t["texto_suave"] else color, True, fuente=fuente)
        elif parte.startswith("`"):
            lz.run(p, parte[1:-1], tam * 0.92, lz.t["primario"], fuente=lz.t["fuente_codigo"])
        else:
            lz.run(p, parte, tam, color, fuente=fuente)


def lista_puntos(lz: Lienzo, x, y, w, h, puntos, tam_max=22, texto_previo: str = ""):
    """Caja de viñetas con tamaño de letra ajustado al espacio disponible."""
    items = _normalizar_puntos(puntos)
    t = lz.t
    tam = tam_max
    for tam in (tam_max, 20, 18, 16, 15, 14, 13, 12):
        alto = 0.0
        if texto_previo:
            alto += _lineas_estimadas(texto_previo, w, tam) * tam * 1.25 / 72 + 0.12
        for nivel, txt in items:
            tn = tam if nivel == 0 else tam - 2
            alto += _lineas_estimadas(txt, w - 0.35 - 0.35 * nivel, tn) * tn * 1.2 / 72 + tn * 0.55 / 72
        if alto <= h:
            break
    tf = lz.texto(x, y, w, h)
    primero = True
    if texto_previo:
        p = tf.paragraphs[0]
        _sin_vineta(p)
        p.line_spacing = 1.15
        p.space_after = Pt(tam * 0.6)
        _escribir_markup(lz, p, texto_previo, tam, t["texto"])
        primero = False
    for nivel, txt in items:
        p = tf.paragraphs[0] if primero else tf.add_paragraph()
        primero = False
        p.line_spacing = 1.08
        if nivel == 0:
            p.space_before = Pt(tam * 0.45)
            _vineta(p, "■", t["primario"], 0.32, 0.32, 70)
            _escribir_markup(lz, p, txt, tam, t["texto"])
        else:
            p.space_before = Pt(tam * 0.2)
            _vineta(p, "–", t["texto_suave"], 0.68, 0.28, 100)
            _escribir_markup(lz, p, txt, tam - 2, t["texto_suave"])
    return tf


# ============================================================================
#  Generador
# ============================================================================

class Generador:
    def __init__(self, ruta_yaml: Path, tema_cli: str | None):
        self.carpeta = ruta_yaml.parent
        self.cfg = leer_yaml(ruta_yaml)
        temas = leer_yaml(TEMAS)
        fuentes = temas.pop("_fuentes", {})
        nombre = tema_cli or self.cfg.get("tema") or "unsa"
        if nombre not in temas:
            raise ErrorDatos(f"Tema '{nombre}' no existe. Disponibles: {', '.join(temas)}")
        self.tema = combinar(combinar(fuentes, temas[nombre]), self.cfg.get("colores") or {})
        self.op = self.cfg.get("opciones") or {}
        # Institución: configuración común + cambios propios de la presentación.
        inst = leer_yaml(CONFIG_INSTITUCION).get("institucion", {})
        self.inst = combinar(inst, self.cfg.get("institucion") or {})
        # Datos de portada: propios, o tomados de otro documento con `desde:`.
        base = _datos_desde(self.carpeta / self.cfg["desde"]) if self.cfg.get("desde") else {}
        self.meta = combinar(base, self.cfg.get("datos") or {})
        self.etq = combinar({"autores": "AUTORES", "docentes": "DOCENTE", "fecha": "FECHA"},
                            self.meta.get("etiquetas") or {})
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(ANCHO), Inches(ALTO)
        self.layout = self.prs.slide_layouts[6]  # en blanco
        self.figura = 0
        self.tabla_n = 0
        self.seccion_actual = ""
        self.num_seccion = 0

    # --- utilidades -----------------------------------------------------------
    def ruta(self, archivo: str) -> Path:
        p = Path(archivo)
        return p if p.is_absolute() else self.carpeta / p

    def ruta_raiz(self, archivo: str) -> Path:
        return RAIZ / archivo

    def nueva(self) -> Lienzo:
        s = self.prs.slides.add_slide(self.layout)
        lz = Lienzo(s, self.tema)
        fondo = s.background.fill
        fondo.solid()
        fondo.fore_color.rgb = _rgb(self.tema["fondo"])
        return lz

    def pie_texto(self) -> str:
        pie = self.op.get("pie", "")
        if pie:
            return pie
        if self.meta.get("pie"):
            return self.meta["pie"]
        partes = [self.meta.get("subtitulo", ""), _titulo_es(self.inst.get("universidad", ""))]
        return "  ·  ".join(p for p in partes if p)

    def marco(self, lz: Lienzo, titulo: str, indice: int, total: int):
        """Título, línea de acento, pie, numeración y barra de progreso."""
        t = self.tema
        if self.seccion_actual and _verdad(self.op.get("mostrar_seccion"), True):
            lz.texto(MX, 0.38, 9.5, 0.3, self.seccion_actual.upper(), 11, t["primario"], True, espaciado=1.5)
        lz.texto(MX, 0.62, ANCHO - 2 * MX - 1.4, 0.85, titulo, 30, t["texto"], True,
                 fuente=t["fuente_titulos"], ancla=MSO_ANCHOR.TOP, interlinea=0.95)
        lz.rect(MX, 1.42, 0.85, 0.06, t["primario"])
        # logo pequeño opcional
        logo = self.inst.get("logo_izquierdo", "")
        if logo and _verdad(self.op.get("logo_en_diapositivas"), True) and t["fondo"].upper() == "#FFFFFF":
            lz.imagen(self.ruta_raiz(logo), ANCHO - MX - 1.5, 0.38, 1.5, 0.62)
        # pie
        lz.texto(MX, ALTO - 0.48, 9, 0.25, self.pie_texto(), 10, t["texto_suave"])
        if _verdad(self.op.get("numeracion"), True):
            lz.texto(ANCHO - MX - 1.5, ALTO - 0.48, 1.5, 0.25, f"{indice} / {total}", 10,
                     t["texto_suave"], True, alinear=PP_ALIGN.RIGHT)
        if _verdad(self.op.get("barra_progreso"), True):
            lz.rect(0, ALTO - 0.07, ANCHO, 0.07, t["fondo_alt"])
            lz.rect(0, ALTO - 0.07, ANCHO * indice / total, 0.07, t["primario"])

    def notas(self, lz: Lienzo, d: dict):
        if d.get("notas"):
            lz.s.notes_slide.notes_text_frame.text = d["notas"]

    # --- tipos de diapositiva ---------------------------------------------------
    def portada(self, d, i, n):
        t, m = self.tema, self.meta
        lz = self.nueva()
        lz.rect(0, 0, ANCHO, 0.14, t["primario"])
        # logos
        li, ld = self.inst.get("logo_izquierdo", ""), self.inst.get("logo_derecho", "")
        if li:
            lz.imagen(self.ruta_raiz(li), MX, 0.45, 2.6, 0.95)
        if ld:
            lz.imagen(self.ruta_raiz(ld), ANCHO - MX - 1.0, 0.45, 1.0, 0.95)
        kicker = d.get("etiqueta") or m.get("etiqueta", "")
        lz.texto(MX, 1.95, 11, 0.35, kicker.upper(), 13, t["primario"], True, espaciado=2)
        titulo = d.get("titulo") or m.get("titulo", "")
        tam = 40 if len(titulo) <= 45 else 34 if len(titulo) <= 80 else 28
        lz.texto(MX, 2.38, 11.6, 1.75, titulo, tam, t["texto"], True, fuente=t["fuente_titulos"],
                 ancla=MSO_ANCHOR.MIDDLE, interlinea=0.95)
        lz.rect(MX, 4.25, 1.1, 0.07, t["primario"])
        sub = d.get("subtitulo") or m.get("subtitulo", "")
        lz.texto(MX, 4.45, 11.6, 0.5, sub, 20, t["texto_suave"])
        # franja inferior con autores, docente y fecha
        y0 = 5.35
        lz.rect(0, y0, ANCHO, ALTO - y0, t["oscuro"])
        cols = [
            (self.etq["autores"], a_lista(d.get("autores") or m.get("autores"))),
            (self.etq["docentes"], a_lista(d.get("docentes") or m.get("docentes"))),
            (self.etq["fecha"], a_lista(d.get("fecha") or m.get("fecha"))),
        ]
        anchos = [5.6, 4.0, 2.2]
        x = MX
        for (etq, vals), w in zip(cols, anchos):
            if not any(vals):
                x += w
                continue
            lz.texto(x, y0 + 0.35, w - 0.2, 0.3, etq, 10.5, t["acento"], True, espaciado=1.5)
            tf = lz.texto(x, y0 + 0.68, w - 0.25, 1.35)
            tam = 15 if len(vals) <= 2 else 13 if len(vals) <= 3 else 11.5
            for k, v in enumerate(vals):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.space_after = Pt(2)
                lz.run(p, v, tam, t["texto_claro"])
            x += w
        self.notas(lz, d)

    def agenda(self, d, i, n):
        t = self.tema
        lz = self.nueva()
        self.marco(lz, d.get("titulo", "Contenido"), i, n)
        puntos = a_lista(d.get("puntos")) or [s.get("titulo", "") for s in self.secciones]
        filas = len(puntos)
        cols = 2 if filas > 5 else 1
        por_col = math.ceil(filas / cols)
        alto_item = min(0.85, (FONDO_CUERPO - TOPE_CUERPO - 0.1) / por_col)
        ancho_col = (ANCHO - 2 * MX - 0.5 * (cols - 1)) / cols
        for k, p in enumerate(puntos):
            c, r = divmod(k, por_col)
            x = MX + c * (ancho_col + 0.5)
            y = TOPE_CUERPO + 0.15 + r * alto_item
            caja = lz.rect(x, y, 0.62, alto_item - 0.17, t["primario"])
            tf = caja.text_frame
            tf.margin_left = tf.margin_right = 0
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.paragraphs[0].alignment = PP_ALIGN.CENTER
            lz.run(tf.paragraphs[0], f"{k + 1:02d}", 16, t["texto_claro"], True)
            lz.texto(x + 0.85, y, ancho_col - 0.9, alto_item - 0.17, p, 20, t["texto"], ancla=MSO_ANCHOR.MIDDLE)
            lz.rect(x + 0.85, y + alto_item - 0.12, ancho_col - 0.9, 0.012, t["fondo_alt"] if t["fondo"] != t["fondo_alt"] else t["texto_suave"])
        self.notas(lz, d)

    def seccion(self, d, i, n):
        t = self.tema
        self.num_seccion += 1
        self.seccion_actual = d.get("titulo", "")
        lz = self.nueva()
        lz.rect(0, 0, ANCHO, ALTO, t["oscuro"])
        lz.rect(0, 0, 0.18, ALTO, t["primario"])
        lz.texto(1.1, 1.7, 4, 1.6, f"{self.num_seccion:02d}", 96, t["primario"], True,
                 fuente=t["fuente_titulos"], ancla=MSO_ANCHOR.BOTTOM)
        lz.rect(1.15, 3.55, 1.2, 0.07, t["acento"])
        lz.texto(1.1, 3.85, 11, 1.3, self.seccion_actual, 40, t["texto_claro"], True,
                 fuente=t["fuente_titulos"], interlinea=0.95)
        if d.get("subtitulo"):
            lz.texto(1.1, 5.2, 10.5, 0.9, d["subtitulo"], 18, _mezclar(t["texto_claro"], t["oscuro"], 0.35))
        if _verdad(self.op.get("barra_progreso"), True):
            lz.rect(0, ALTO - 0.07, ANCHO * i / n, 0.07, t["primario"])
        self.notas(lz, d)

    def contenido(self, d, i, n):
        lz = self.nueva()
        self.marco(lz, d.get("titulo", ""), i, n)
        alto = FONDO_CUERPO - TOPE_CUERPO
        if d.get("destacado"):
            alto -= 1.05
        lista_puntos(lz, MX, TOPE_CUERPO, ANCHO - 2 * MX, alto, d.get("puntos"), texto_previo=d.get("texto", ""))
        if d.get("destacado"):
            self._franja_destacada(lz, d["destacado"], FONDO_CUERPO - 0.9)
        self.notas(lz, d)

    def _franja_destacada(self, lz, texto, y):
        t = self.tema
        lz.rect(MX, y, ANCHO - 2 * MX, 0.85, t["fondo_alt"])
        lz.rect(MX, y, 0.08, 0.85, t["primario"])
        tf = lz.texto(MX + 0.3, y, ANCHO - 2 * MX - 0.5, 0.85, ancla=MSO_ANCHOR.MIDDLE)
        _escribir_markup(lz, tf.paragraphs[0], texto, 17, t["texto"])

    def _columna(self, lz, col: dict, x, y, w, h):
        t = self.tema
        if col.get("subtitulo"):
            lz.texto(x, y, w, 0.4, col["subtitulo"], 19, t["primario"], True)
            lz.rect(x, y + 0.45, w, 0.015, t["texto_suave"])
            y, h = y + 0.62, h - 0.62
        if col.get("imagen"):
            self._imagen_con_pie(lz, col["imagen"], col.get("pie", ""), x, y, w, h)
        else:
            lista_puntos(lz, x, y, w, h, col.get("puntos"), tam_max=20, texto_previo=col.get("texto", ""))

    def dos_columnas(self, d, i, n):
        lz = self.nueva()
        self.marco(lz, d.get("titulo", ""), i, n)
        sep = 0.6
        w = (ANCHO - 2 * MX - sep) / 2
        h = FONDO_CUERPO - TOPE_CUERPO
        self._columna(lz, d.get("izquierda") or {}, MX, TOPE_CUERPO, w, h)
        self._columna(lz, d.get("derecha") or {}, MX + w + sep, TOPE_CUERPO, w, h)
        self.notas(lz, d)

    def _imagen_con_pie(self, lz, archivo, pie, x, y, w, h):
        t = self.tema
        alto_pie = 0.45 if pie else 0
        _, (px, py, pw, ph) = lz.imagen(self.ruta(archivo), x, y, w, h - alto_pie, borde=_mezclar(t["texto_suave"], t["fondo"], 0.5))
        if pie:
            self.figura += 1
            tf = lz.texto(x, py + ph + 0.1, w, 0.35, alinear=PP_ALIGN.CENTER)
            lz.run(tf.paragraphs[0], f"Figura {self.figura}. ", 12, t["texto"], True)
            lz.run(tf.paragraphs[0], pie, 12, t["texto_suave"])

    def imagen(self, d, i, n):
        lz = self.nueva()
        self.marco(lz, d.get("titulo", ""), i, n)
        h = FONDO_CUERPO - TOPE_CUERPO
        if d.get("puntos") or d.get("texto"):
            wt = 4.6
            lista_puntos(lz, MX, TOPE_CUERPO, wt, h, d.get("puntos"), tam_max=20, texto_previo=d.get("texto", ""))
            self._imagen_con_pie(lz, d["imagen"], d.get("pie", ""), MX + wt + 0.5, TOPE_CUERPO, ANCHO - 2 * MX - wt - 0.5, h)
        else:
            self._imagen_con_pie(lz, d["imagen"], d.get("pie", ""), MX, TOPE_CUERPO, ANCHO - 2 * MX, h)
        self.notas(lz, d)

    def tabla(self, d, i, n):
        t = self.tema
        lz = self.nueva()
        self.marco(lz, d.get("titulo", ""), i, n)
        enc = a_lista(d.get("encabezado"))
        filas = [a_lista(f) for f in a_lista(d.get("filas"))]
        ncol = max([len(enc)] + [len(f) for f in filas])
        nfil = len(filas) + (1 if enc else 0)
        pie = d.get("pie", "")
        h_disp = FONDO_CUERPO - TOPE_CUERPO - (0.5 if pie else 0) - 0.1
        alto_fila = min(0.55, h_disp / max(nfil, 1))
        tam = 16 if alto_fila >= 0.5 else 14 if alto_fila >= 0.4 else 12
        w = ANCHO - 2 * MX
        y = TOPE_CUERPO + 0.1
        if pie:
            self.tabla_n += 1
            tf = lz.texto(MX, y, w, 0.35, alinear=PP_ALIGN.LEFT)
            lz.run(tf.paragraphs[0], f"Tabla {self.tabla_n}. ", 13, t["texto"], True)
            lz.run(tf.paragraphs[0], pie, 13, t["texto_suave"])
            y += 0.5
        forma = lz.s.shapes.add_table(nfil, ncol, Inches(MX), Inches(y), Inches(w), Inches(alto_fila * nfil))
        tb = forma.table
        # quitar el estilo por defecto para controlar todos los colores
        tblPr = tb._tbl.tblPr
        tblPr.set("bandRow", "0")
        tblPr.set("firstRow", "0")
        anchos = a_lista(d.get("anchos"))
        if anchos and len(anchos) == ncol:
            total = sum(_num(a, 1) for a in anchos)
            for k, a in enumerate(anchos):
                tb.columns[k].width = Emu(int(Inches(w) * _num(a, 1) / total))
        todas = ([enc] if enc else []) + filas
        for r, fila in enumerate(todas):
            tb.rows[r].height = Inches(alto_fila)
            es_enc = bool(enc) and r == 0
            for c in range(ncol):
                celda = tb.cell(r, c)
                celda.fill.solid()
                if es_enc:
                    celda.fill.fore_color.rgb = _rgb(t["primario"])
                else:
                    zebra = (r - (1 if enc else 0)) % 2 == 1
                    celda.fill.fore_color.rgb = _rgb(t["fondo_alt"] if zebra else t["fondo"])
                celda.margin_left = celda.margin_right = Inches(0.12)
                celda.margin_top = celda.margin_bottom = Inches(0.04)
                celda.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf = celda.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                valor = fila[c] if c < len(fila) else ""
                lz.run(p, valor, tam, t["texto_claro"] if es_enc else t["texto"], es_enc)
                p.alignment = PP_ALIGN.CENTER if es_enc else PP_ALIGN.LEFT
        # línea inferior de cierre
        lz.rect(MX, y + alto_fila * nfil, w, 0.03, t["primario"])
        self.notas(lz, d)

    def codigo(self, d, i, n):
        t = self.tema
        lz = self.nueva()
        self.marco(lz, d.get("titulo", ""), i, n)
        if d.get("archivo"):
            ruta = self.ruta(d["archivo"])
            if not ruta.exists():
                raise ErrorDatos(f"No existe el archivo de código: {ruta}")
            codigo = ruta.read_text(encoding="utf-8").rstrip()
            nombre = d.get("nombre") or ruta.name
        else:
            codigo = str(d.get("codigo", "")).rstrip()
            nombre = d.get("nombre", "")
        lineas = codigo.splitlines() or [""]
        explicacion = d.get("puntos") or d.get("texto")
        x, w = MX, ANCHO - 2 * MX
        if explicacion:
            wt = 4.3
            lista_puntos(lz, MX, TOPE_CUERPO, wt, FONDO_CUERPO - TOPE_CUERPO, d.get("puntos"),
                         tam_max=18, texto_previo=d.get("texto", ""))
            x, w = MX + wt + 0.4, ANCHO - 2 * MX - wt - 0.4
        h = FONDO_CUERPO - TOPE_CUERPO
        fondo_cod, barra_cod = "#1E1E1E", "#2D2D30"
        lz.rect(x, TOPE_CUERPO, w, 0.4, barra_cod)
        for k, c in enumerate(("#FF5F57", "#FEBC2E", "#28C840")):
            lz.rect(x + 0.18 + k * 0.22, TOPE_CUERPO + 0.13, 0.14, 0.14, c, MSO_SHAPE.OVAL)
        if nombre:
            lz.texto(x + 1.0, TOPE_CUERPO + 0.06, w - 1.2, 0.3, nombre, 11, "#CCCCCC", fuente=t["fuente_codigo"],
                     ancla=MSO_ANCHOR.MIDDLE)
        alto_util = h - 0.4 - 0.3
        ancho_max = max(len(l.expandtabs(4)) for l in lineas) + len(str(len(lineas))) + 3
        # Consolas: ~0.55 em por carácter; se deja margen para que no se partan líneas
        tam = min(16, alto_util * 72 / (len(lineas) * 1.3), (w - 0.45) * 72 / (ancho_max * 0.58))
        tam = max(8, math.floor(tam * 2) / 2)
        # la caja se ajusta al código si este es corto
        h_caja = min(h - 0.4, len(lineas) * tam * 1.3 / 72 + 0.4)
        lz.rect(x, TOPE_CUERPO + 0.4, w, h_caja, fondo_cod)
        tf = lz.texto(x + 0.2, TOPE_CUERPO + 0.55, w - 0.3, h_caja - 0.2)
        tokens = resaltar(codigo, d.get("lenguaje", ""), nombre)
        ancho_num = len(str(len(lineas)))
        for k, linea in enumerate(tokens):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.line_spacing = 1.05
            lz.run(p, f"{k + 1:>{ancho_num}}  ", tam, "#6E7681", fuente=t["fuente_codigo"])
            for texto, color, cursiva in linea:
                lz.run(p, texto.expandtabs(4), tam, color, cursiva=cursiva, fuente=t["fuente_codigo"])
        self.notas(lz, d)

    def destacado(self, d, i, n):
        t = self.tema
        lz = self.nueva()
        if d.get("titulo"):
            self.marco(lz, d["titulo"], i, n)
            y0 = TOPE_CUERPO + 0.3
        else:
            y0 = 1.2
            if _verdad(self.op.get("barra_progreso"), True):
                lz.rect(0, ALTO - 0.07, ANCHO, 0.07, t["fondo_alt"])
                lz.rect(0, ALTO - 0.07, ANCHO * i / n, 0.07, t["primario"])
        lz.texto(1.3, y0 - 0.35, 1.5, 1.5, "“", 120, t["primario"], True, fuente="Georgia")
        texto = d.get("texto", "")
        tam = 30 if len(texto) < 120 else 26 if len(texto) < 220 else 22
        lz.texto(2.3, y0 + 0.4, ANCHO - 4.3, 3.2, texto, tam, t["texto"], cursiva=True,
                 fuente=t["fuente_titulos"], ancla=MSO_ANCHOR.MIDDLE, interlinea=1.15)
        if d.get("autor"):
            lz.rect(2.3, y0 + 3.85, 0.6, 0.04, t["primario"])
            lz.texto(3.05, y0 + 3.7, 8, 0.4, d["autor"], 16, t["texto_suave"], True)
        self.notas(lz, d)

    def metricas(self, d, i, n):
        t = self.tema
        lz = self.nueva()
        self.marco(lz, d.get("titulo", ""), i, n)
        items = a_lista(d.get("items"))[:4]
        if not items:
            return
        sep = 0.35
        w = (ANCHO - 2 * MX - sep * (len(items) - 1)) / len(items)
        y, h = TOPE_CUERPO + 0.35, 2.9
        for k, it in enumerate(items):
            x = MX + k * (w + sep)
            lz.rect(x, y, w, h, t["fondo_alt"])
            lz.rect(x, y, w, 0.08, t["primario"])
            lz.texto(x + 0.3, y + 0.45, w - 0.6, 1.2, it.get("valor", ""), 48, t["primario"], True,
                     fuente=t["fuente_titulos"], ancla=MSO_ANCHOR.MIDDLE)
            lz.texto(x + 0.3, y + 1.7, w - 0.6, 0.45, it.get("etiqueta", ""), 17, t["texto"], True)
            if it.get("detalle"):
                lz.texto(x + 0.3, y + 2.15, w - 0.6, 0.65, it["detalle"], 13, t["texto_suave"])
        if d.get("texto"):
            lista_puntos(lz, MX, y + h + 0.35, ANCHO - 2 * MX, FONDO_CUERPO - y - h - 0.35, [], 18, d["texto"])
        self.notas(lz, d)

    def referencias(self, d, i, n):
        t = self.tema
        lz = self.nueva()
        self.marco(lz, d.get("titulo", "Referencias"), i, n)
        archivo = d.get("archivo") or self.meta.get("referencias", "")
        if not archivo:
            raise ErrorDatos("La diapositiva de referencias necesita un archivo .bib")
        entradas = leer_bib(self.ruta(archivo))
        claves = a_lista(d.get("claves"))
        if claves:
            por_clave = {e["clave"]: e for e in entradas}
            entradas = [por_clave[c] for c in claves if c in por_clave]
        refs = [formato_ieee(e) for e in entradas]
        h = FONDO_CUERPO - TOPE_CUERPO
        w = ANCHO - 2 * MX
        for tam in (16, 15, 14, 13, 12, 11, 10):
            alto = sum(_lineas_estimadas("".join(s for s, _ in r), w - 0.6, tam) * tam * 1.2 / 72 + tam * 0.6 / 72
                       for r in refs)
            if alto <= h:
                break
        tf = lz.texto(MX, TOPE_CUERPO, w, h)
        for k, segmentos in enumerate(refs):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.space_before = Pt(tam * 0.6)
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(Inches(0.6)))
            pPr.set("indent", str(-Inches(0.6)))
            lz.run(p, f"[{k + 1}]\t", tam, t["primario"], True)
            for texto, cursiva in segmentos:
                lz.run(p, texto, tam, t["texto"], cursiva=cursiva)
        self.notas(lz, d)

    def cierre(self, d, i, n):
        t = self.tema
        lz = self.nueva()
        lz.rect(0, 0, ANCHO, ALTO, t["oscuro"])
        lz.rect(0, 0, ANCHO, 0.14, t["primario"])
        lz.texto(MX, 2.1, ANCHO - 2 * MX, 1.3, d.get("titulo", "¡Gracias!"), 60, t["texto_claro"], True,
                 fuente=t["fuente_titulos"], alinear=PP_ALIGN.CENTER, ancla=MSO_ANCHOR.MIDDLE)
        lz.rect(ANCHO / 2 - 0.6, 3.6, 1.2, 0.07, t["primario"])
        lz.texto(MX, 3.9, ANCHO - 2 * MX, 0.6, d.get("subtitulo", "¿Preguntas?"), 24, t["acento"],
                 alinear=PP_ALIGN.CENTER)
        contacto = d.get("contacto") or "  ·  ".join(a_lista(self.meta.get("autores")))
        if contacto:
            lz.texto(1.2, 5.4, ANCHO - 2.4, 0.9, contacto, 13, _mezclar(t["texto_claro"], t["oscuro"], 0.3),
                     alinear=PP_ALIGN.CENTER)
        lz.texto(MX, ALTO - 0.55, ANCHO - 2 * MX, 0.3, self.pie_texto(), 10,
                 _mezclar(t["texto_claro"], t["oscuro"], 0.45), alinear=PP_ALIGN.CENTER)
        self.notas(lz, d)

    # --- ejecución ----------------------------------------------------------------
    def construir(self) -> Presentation:
        diapositivas = a_lista(self.cfg.get("diapositivas"))
        if not diapositivas:
            raise ErrorDatos(f"{self.carpeta.name}: el YAML no tiene ninguna diapositiva en 'diapositivas:'")
        self.secciones = [d for d in diapositivas if d.get("tipo") == "seccion"]
        n = len(diapositivas)
        for i, d in enumerate(diapositivas, start=1):
            tipo = d.get("tipo", "contenido")
            if tipo not in TIPOS:
                raise ErrorDatos(f"Diapositiva {i}: tipo desconocido '{tipo}'. Tipos: {', '.join(TIPOS)}")
            getattr(self, tipo)(d, i, n)
        self.prs.core_properties.title = self.meta.get("titulo", "")
        self.prs.core_properties.author = ", ".join(a_lista(self.meta.get("autores")))
        return self.prs


TIPOS = ["portada", "agenda", "seccion", "contenido", "dos_columnas", "imagen", "tabla",
         "codigo", "destacado", "metricas", "referencias", "cierre"]

def _datos_desde(ruta: Path) -> dict:
    """Datos de portada tomados de otro documento (por ahora, un informe de laboratorio)."""
    d = leer_yaml(ruta)
    info = d.get("informacion")
    if not info:
        raise ErrorDatos(f"`desde: {ruta.name}`: el archivo no tiene la sección 'informacion'")
    numero = info.get("numero", "")
    tipo = d.get("documento", {}).get("tipo", "")
    return {
        "titulo": info.get("titulo", ""),
        "subtitulo": info.get("asignatura", ""),
        "etiqueta": " · ".join(x for x in (tipo, f"N.º {numero}" if numero else "") if x),
        "autores": a_lista(info.get("integrantes")),
        "docentes": a_lista(info.get("docentes")),
        "fecha": " · ".join(x for x in (info.get("fecha", ""), info.get("anio_lectivo", "")) if x),
        "referencias": d.get("referencias", {}).get("archivo", ""),
        "etiquetas": {"autores": "INTEGRANTES"},
    }


def nombre_pptx(ruta_yaml: Path, cfg: dict, meta: dict) -> str:
    archivo = cfg.get("archivo", "")
    if archivo:
        return archivo.removesuffix(".pptx") + ".pptx"
    titulo = meta.get("titulo", "") or ruta_yaml.parent.name
    return f"Presentacion_{slug(titulo)}.pptx"


def generar(ruta_yaml: Path, salida: Path, opciones: dict) -> list[Path]:
    g = Generador(ruta_yaml, opciones.get("tema"))
    prs = g.construir()
    salida.mkdir(parents=True, exist_ok=True)
    nombre = nombre_pptx(ruta_yaml, g.cfg, g.meta)
    salidas = opciones["salidas"]
    # El PDF se exporta desde el PPTX; si no se pidió el PPTX, se usa uno temporal.
    if "pptx" in salidas:
        pptx = salida / nombre
    else:
        (RAIZ / ".build").mkdir(exist_ok=True)
        pptx = RAIZ / ".build" / nombre
    try:
        prs.save(pptx)
    except PermissionError as e:
        raise ErrorDatos(f"No se pudo guardar {pptx.name}: ciérralo en PowerPoint y vuelve a intentar.") from e
    archivos = [pptx] if "pptx" in salidas else []
    if "pdf" in salidas:
        from office import a_pdf
        archivos.append(a_pdf(pptx, salida / (Path(nombre).stem + ".pdf")))
    return archivos


def listar_temas() -> None:
    temas = leer_yaml(TEMAS)
    temas.pop("_fuentes", None)
    print("Temas disponibles:")
    for clave, t in temas.items():
        print(f"  {clave:<10} {t.get('nombre', '')}  (primario {t.get('primario')})")
