"""Versión Word (.docx) del informe de laboratorio.

Parte de la plantilla oficial (plantillas/informe-laboratorio/plantilla.docx):
copia sus filas, barras, numeración y encabezado reales, de modo que el
resultado conserva exactamente el formato de Word. El contenido de las
secciones (.typ o texto del YAML) se escribe con docx_render.
"""

from __future__ import annotations

import io
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Twips
from docx.table import Table, _Cell

from bib import formato_ieee
from datos import RAIZ, ErrorDatos, a_lista
from docx_render import Renderizador, espaciado, poner_texto, rpr, rpr_de, run, vaciar

PLANTILLA_DOCX = RAIZ / "plantillas" / "informe-laboratorio" / "plantilla.docx"

ANCHO_TABLA = 10266            # twips (igual que la plantilla)
MARGEN_CELDA = 108
SANGRIA_APARTADO = 322         # texto de los apartados (I., II., ...)
SANGRIA_PREGUNTA = 740         # viñetas del cuestionario


class ConstructorDocx:
    def __init__(self, datos: dict, carpeta: Path):
        self.d = datos
        estilo = datos.get("estilo", {})
        self.color_datos = estilo.get("color_datos", "#2F5496")
        self.ref = datos.get("referencias", {}) or {}
        self.r = Renderizador(carpeta, {
            "fuente_codigo": estilo.get("fuente_codigo", "Consolas"),
            "color_enc": estilo.get("color_barra", "#C8310E"),
            "color_texto_enc": estilo.get("color_texto_barra", "#FFFFFF"),
            "nota_borde": estilo.get("color_barra", "#C8310E"),
        }, self.ref.get("archivo", ""), self.ref.get("incluir_todas") == "true")

    def bloques(self, celda: _Cell, bloques: list[dict], sangria: int):
        self.r.bloques(celda, bloques, sangria, ANCHO_TABLA - 2 * MARGEN_CELDA - sangria)

    # --- estructura del documento ---------------------------------------------------
    def construir(self, destino: Path):
        doc = Document(str(PLANTILLA_DOCX))
        self.doc = doc
        info = self.d.get("informacion", {})
        et = self.d.get("etiquetas", {}) or {}
        etq = lambda k, defecto: et.get(k) or defecto  # noqa: E731

        self._encabezado(doc)

        body = doc.element.body
        sectpr = body.find(qn("w:sectPr"))
        elementos = [el for el in body if el.tag != qn("w:sectPr")]
        tablas = [el for el in elementos if el.tag == qn("w:tbl")]
        if len(tablas) < 4:
            raise ErrorDatos("plantilla.docx no tiene la estructura esperada (4 tablas)")
        p_titulo, p_sep = deepcopy(elementos[0]), deepcopy(elementos[2])
        t_info, t_sec, t_retro, t_refs = (deepcopy(t) for t in tablas[:4])
        filas_sec = t_sec.findall(qn("w:tr"))
        proto_barra, proto_apartado = filas_sec[0], filas_sec[1]
        proto_pregunta = filas_sec[2].find(qn("w:tc")).findall(qn("w:p"))[1]
        proto_caja = t_retro.findall(qn("w:tr"))[1]
        self.abstract_romano = self._abstract_de(proto_apartado)
        for el in elementos:
            body.remove(el)
        agregar = sectpr.addprevious

        # Título
        poner_texto(p_titulo, self.d.get("documento", {}).get("tipo", "INFORME DE LABORATORIO"))
        agregar(p_titulo)

        # Información básica
        self._info(t_info, info, etq)
        agregar(t_info)
        agregar(deepcopy(p_sep))

        # Contenido: primero se carga y numera todo (para citas y referencias cruzadas)
        plan = []
        for s in a_lista(self.d.get("secciones")):
            apartados = []
            for a in a_lista(s.get("apartados")):
                preguntas = [(q.get("pregunta", "") if isinstance(q, dict) else q,
                              self.r.cargar(q.get("respuesta")) if isinstance(q, dict) else [])
                             for q in a_lista(a.get("preguntas"))]
                apartados.append((a, preguntas, self.r.cargar(a.get("contenido"))))
            plan.append((s, apartados, self.r.cargar(s.get("contenido"))))
        for s, apartados, contenido in plan:
            for a, preguntas, bloques in apartados:
                for _, respuesta in preguntas:
                    self.r.numerar(respuesta)
                self.r.numerar(bloques)
            self.r.numerar(contenido)
        self.r.cerrar_citas()

        for s, apartados, contenido in plan:
            agregar(self._seccion(t_sec, proto_barra, proto_apartado, proto_pregunta, proto_caja,
                                  s, apartados, contenido))
            agregar(deepcopy(p_sep))

        # Referencias
        agregar(self._referencias(t_refs))
        agregar(deepcopy(p_sep))

        doc.core_properties.title = info.get("titulo", "")
        doc.core_properties.author = ", ".join(a_lista(info.get("integrantes")))
        try:
            doc.save(str(destino))
        except PermissionError as e:
            raise ErrorDatos(f"No se pudo guardar {destino.name}: ciérralo en Word y vuelve a intentar.") from e

    # --- encabezado -------------------------------------------------------------
    def _encabezado(self, doc):
        inst = self.d.get("institucion", {})
        enc = self.d.get("encabezado", {})
        header = doc.sections[0].header
        t = header.tables[0]
        for p, clave in zip(t.cell(0, 1).paragraphs, ("universidad", "facultad", "escuela")):
            poner_texto(p._p, inst.get(clave, ""))
        p = t.cell(1, 0).paragraphs[0]._p
        runs = p.findall(qn("w:r"))
        rpr_b, rpr_n = deepcopy(runs[0].find(qn("w:rPr"))), deepcopy(runs[-1].find(qn("w:rPr")))
        vaciar(p)
        run(p, "Formato: ", rpr_b)
        run(p, enc.get("formato", ""), rpr_n)
        poner_texto(t.cell(2, 0).paragraphs[0]._p, f"Aprobación:  {enc.get('aprobacion', '')}")
        poner_texto(t.cell(2, 1).paragraphs[0]._p, f"Código: {enc.get('codigo', '')}")
        # Logos: se reemplaza la imagen de la plantilla por la configurada
        for blip, clave in zip(header._element.iter(qn("a:blip")), ("logo_izquierdo", "logo_derecho")):
            ruta = RAIZ / inst.get(clave, "") if inst.get(clave) else None
            if ruta and ruta.exists():
                from PIL import Image
                buf = io.BytesIO()
                with Image.open(ruta) as im:
                    im.save(buf, "PNG")
                header.part.related_parts[blip.get(qn("r:embed"))]._blob = buf.getvalue()

    # --- información básica -------------------------------------------------------
    def _valor(self, celda: _Cell, texto: str):
        p = celda.paragraphs[0]._p
        vaciar(p)
        if texto:
            run(p, texto, rpr(cursiva=True, color=self.color_datos))

    def _info(self, t_el, info: dict, etq):
        filas = Table(t_el, self.doc._body).rows
        poner_texto(filas[0].cells[0].paragraphs[0]._p, etq("informacion_basica", "INFORMACIÓN BÁSICA"))
        pares = [
            (1, 0, "asignatura", "ASIGNATURA: "), (2, 0, "titulo", "TÍTULO DE LA PRÁCTICA: "),
            (3, 0, "numero", "NÚMERO DE PRÁCTICA:"), (3, 2, "anio_lectivo", "AÑO LECTIVO:"),
            (3, 4, "semestre", "NRO. SEMESTRE:"), (4, 0, "fecha", "FECHA DE PRESENTACIÓN"),
            (4, 2, "hora", "HORA DE PRESENTACIÓN"), (5, 4, "nota", "NOTA:"),
        ]
        for r, c, clave, defecto in pares:
            poner_texto(filas[r].cells[c].paragraphs[0]._p, etq(clave, defecto))
            self._valor(filas[r].cells[c + 1], info.get(clave, ""))
        # Integrantes: etiqueta + una línea por nombre (mínimo dos, como la plantilla)
        cel = filas[5].cells[0]
        poner_texto(cel.paragraphs[0]._p, etq("integrantes", "INTEGRANTE (s): "))
        lineas = cel.paragraphs[1:]
        nombres = [str(x) for x in a_lista(info.get("integrantes")) if x]
        while len(lineas) < len(nombres):
            lineas[-1]._p.addnext(deepcopy(lineas[-1]._p))
            lineas = cel.paragraphs[1:]
        for k, p in enumerate(lineas):
            vaciar(p._p)
            if k < len(nombres):
                run(p._p, nombres[k], rpr(cursiva=True, color=self.color_datos, tam=11))
        # Docentes
        p = filas[6].cells[0].paragraphs[0]._p
        rpr_b = rpr_de(p)
        vaciar(p)
        run(p, etq("docentes", "DOCENTE(s):") + " ", rpr_b)
        run(p, ", ".join(str(x) for x in a_lista(info.get("docentes")) if x),
            rpr(cursiva=True, color=self.color_datos))

    # --- secciones -------------------------------------------------------------------
    def _abstract_de(self, tr) -> str:
        num_id = tr.find(".//" + qn("w:numId")).get(qn("w:val"))
        for n in self.doc.part.numbering_part.element.findall(qn("w:num")):
            if n.get(qn("w:numId")) == num_id:
                return n.find(qn("w:abstractNumId")).get(qn("w:val"))
        raise ErrorDatos("No se encontró la numeración romana en plantilla.docx")

    def _nueva_numeracion(self) -> str:
        """Una lista I, II, III... que reinicia en cada sección."""
        numbering = self.doc.part.numbering_part.element
        nuevo = str(max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1)
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), nuevo)
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), self.abstract_romano)
        num.append(a)
        ov = OxmlElement("w:lvlOverride")
        ov.set(qn("w:ilvl"), "0")
        st = OxmlElement("w:startOverride")
        st.set(qn("w:val"), "1")
        ov.append(st)
        num.append(ov)
        numbering.append(num)
        return nuevo

    @staticmethod
    def _alto(tr, twips: int):
        trpr = tr.find(qn("w:trPr"))
        if trpr is None:
            trpr = OxmlElement("w:trPr")
            tr.insert(0, trpr)
        for h in trpr.findall(qn("w:trHeight")):
            trpr.remove(h)
        h = OxmlElement("w:trHeight")
        h.set(qn("w:val"), str(twips))
        trpr.append(h)

    @staticmethod
    def _alto_de(valor, defecto: int) -> int:
        try:
            return int(float(valor) * 567) if valor else defecto
        except ValueError:
            return defecto

    def _seccion(self, t_sec, proto_barra, proto_apartado, proto_pregunta, proto_caja, s, apartados, contenido):
        tbl = deepcopy(t_sec)
        for tr in tbl.findall(qn("w:tr")):
            tbl.remove(tr)
        tabla = Table(tbl, self.doc._body)

        barra = deepcopy(proto_barra)
        poner_texto(barra.find(qn("w:tc")).find(qn("w:p")), s.get("titulo", ""))
        tbl.append(barra)

        if apartados:
            num_id = self._nueva_numeracion()
            for k, (a, preguntas, bloques) in enumerate(apartados):
                tr = deepcopy(proto_apartado)
                tc = tr.find(qn("w:tc"))
                ps = tc.findall(qn("w:p"))
                for p in ps[1:]:
                    tc.remove(p)
                titulo = ps[0]
                titulo.find(".//" + qn("w:numId")).set(qn("w:val"), num_id)
                poner_texto(titulo, a.get("titulo", ""))
                self._alto(tr, self._alto_de(a.get("alto_minimo"), 1353 if k == 0 else 1088))
                tbl.append(tr)
                celda = _Cell(tc, tabla)
                for pregunta, respuesta in preguntas:
                    p = deepcopy(proto_pregunta)
                    p.find(".//" + qn("w:numId")).set(qn("w:val"), num_id)
                    vaciar(p)
                    run(p, str(pregunta), rpr(cursiva=True, color=self.color_datos))
                    if respuesta:  # la pregunta no queda separada de su respuesta
                        ppr = p.find(qn("w:pPr"))
                        if ppr.find(qn("w:keepNext")) is None:
                            ppr.insert(1, OxmlElement("w:keepNext"))
                    tc.append(p)
                    if respuesta:
                        self.bloques(celda, respuesta, SANGRIA_PREGUNTA)
                self.bloques(celda, bloques, SANGRIA_APARTADO)
                self._cerrar_celda(tc)
        else:
            tr = deepcopy(proto_caja)
            tc = tr.find(qn("w:tc"))
            if contenido:
                for p in tc.findall(qn("w:p")):
                    tc.remove(p)
                self.bloques(_Cell(tc, tabla), contenido, 0)
            self._cerrar_celda(tc)
            self._alto(tr, self._alto_de(s.get("alto_minimo"), 1088))
            tbl.append(tr)
        return tbl

    @staticmethod
    def _cerrar_celda(tc):
        """Word exige que una celda termine en párrafo."""
        hijos = [h for h in tc if h.tag != qn("w:tcPr")]
        if not hijos or hijos[-1].tag != qn("w:p"):
            tc.append(OxmlElement("w:p"))

    def _referencias(self, tbl):
        filas = tbl.findall(qn("w:tr"))
        poner_texto(filas[0].find(qn("w:tc")).find(qn("w:p")), self.ref.get("titulo", "REFERENCIAS Y BIBLIOGRAFÍA"))
        tc = filas[1].find(qn("w:tc"))
        if self.r.citas:
            for p in tc.findall(qn("w:p")):
                tc.remove(p)
            celda = _Cell(tc, Table(tbl, self.doc._body))
            for k, clave in enumerate(self.r.citas):
                p = celda.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                pf = p.paragraph_format
                pf.left_indent = Twips(567)
                pf.first_line_indent = Twips(-567)
                pf.tab_stops.add_tab_stop(Twips(567))
                espaciado(p, 3 if k == 0 else 0, 6, 1.0)
                run(p._p, f"[{k + 1}]\t", rpr(tam=11))
                for texto, cursiva in formato_ieee(self.r.bib[clave]):
                    run(p._p, texto, rpr(cursiva=cursiva or None, tam=11))
        self._cerrar_celda(tc)
        return tbl


def generar_docx(datos: dict, carpeta: Path, destino: Path) -> Path:
    ConstructorDocx(datos, carpeta).construir(destino)
    return destino
