"""Versión Word (.docx) del informe académico.

Usa como base plantillas/informe-academico/plantilla.docx (estilos y tamaño
de página de la plantilla original, sin contenido). La portada replica la
original; el índice es un campo TOC de Word que se actualiza con Word si está
instalado (si no, Word lo actualiza al abrir el archivo).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

from bib import formato_ieee
from datos import RAIZ, ErrorDatos, a_lista
from docx_render import Renderizador, campo, espaciado, rpr, run

PLANTILLA_DOCX = RAIZ / "plantillas" / "informe-academico" / "plantilla.docx"
ANCHO = 11909 - 2 * 1440          # ancho útil de la página (twips)


def _si(valor, defecto=True) -> bool:
    if valor in (None, ""):
        return defecto
    return str(valor).strip().lower() in ("true", "si", "sí", "1")


class ConstructorAcademico:
    def __init__(self, datos: dict, carpeta: Path):
        self.d = datos
        self.por = datos.get("portada", {}) or {}
        self.inst = datos.get("institucion", {}) or {}
        est = datos.get("estilo", {}) or {}
        self.fuente = est.get("fuente", "Times New Roman")
        self.ref = datos.get("referencias", {}) or {}
        self.r = Renderizador(carpeta, {
            "fuente": self.fuente,
            "fuente_codigo": est.get("fuente_codigo", "Consolas"),
            "tam": 11, "interlineado": 1.15, "despues": 12, "sangria_primera": 720,
            "color_enc": est.get("color_encabezado", "#FCE5CD"),
            "color_texto_enc": est.get("color_texto_encabezado", "#000000"),
            "borde_tabla": (est.get("color_borde", "#C4C7C5"), 5),
            "tam_tabla": 10, "margen_tabla": (120, 180), "alinear_enc": "izquierda",
            "tam_pie": 11, "pie": "apa", "tabla_pie": "izquierda_negrita",
            "nota_borde": est.get("color_acento", "#B45F06"),
            "nota_fondo": "#FDF2E6",
            "estilos_titulo": {1: "Heading 2", 2: "Heading 3", 3: "Heading 3"},
            "campos_seq": True,
            "vinetas": ("●", "○"),
        }, self.ref.get("archivo", ""), _si(self.ref.get("incluir_todas"), False))

    # --- estilos del documento ----------------------------------------------------
    def _estilos(self, doc):
        def fuente_estilo(st, tam, negrita=None, cursiva=None):
            st.font.name = self.fuente
            st.font.size = Pt(tam)
            if negrita is not None:
                st.font.bold = negrita
            if cursiva is not None:
                st.font.italic = cursiva
            st.font.color.rgb = RGBColor(0, 0, 0)
            rfonts = st.element.get_or_add_rPr().get_or_add_rFonts()
            for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
                rfonts.set(qn(a), self.fuente)

        normal = doc.styles.default(WD_STYLE_TYPE.PARAGRAPH)
        fuente_estilo(normal, 11)
        normal.paragraph_format.line_spacing = 1.15
        for nombre, tam, cursiva, antes, despues in (("Heading 1", 12, False, 18, 12),
                                                      ("Heading 2", 11, False, 12, 8),
                                                      ("Heading 3", 11, True, 10, 6)):
            st = doc.styles[nombre]
            fuente_estilo(st, tam, True, cursiva)
            pf = st.paragraph_format
            pf.space_before, pf.space_after = Pt(antes), Pt(despues)
            pf.keep_with_next = True
            pf.line_spacing = 1.15
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        # Estilos del índice (Word los usa al generar el TOC)
        for nivel in (1, 2, 3):
            nombre = f"toc {nivel}"
            try:
                st = doc.styles[nombre]
            except KeyError:
                st = doc.styles.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
                st.base_style = normal
            fuente_estilo(st, 11, nivel == 1)
            pf = st.paragraph_format
            pf.left_indent = Twips(240 * (nivel - 1))
            pf.space_before = Pt(6 if nivel == 1 else 0)
            pf.space_after = Pt(3)
            pf.tab_stops.add_tab_stop(Twips(ANCHO), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        try:
            st = doc.styles["table of figures"]
        except KeyError:
            st = doc.styles.add_style("table of figures", WD_STYLE_TYPE.PARAGRAPH)
            st.base_style = normal
        fuente_estilo(st, 11)
        st.paragraph_format.tab_stops.add_tab_stop(Twips(ANCHO), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    # --- utilidades ------------------------------------------------------------------
    def _p(self, doc, alinear=WD_ALIGN_PARAGRAPH.CENTER, interlineado=1.5):
        p = doc.add_paragraph()
        p.alignment = alinear
        espaciado(p, 0, 0, interlineado)
        return p

    def _salto_pagina(self, doc):
        p = doc.add_paragraph()
        espaciado(p, 0, 0)
        r = OxmlElement("w:r")
        br = OxmlElement("w:br")
        br.set(qn("w:type"), "page")
        r.append(br)
        p._p.append(r)

    def _titulo_centrado(self, doc, texto):
        p = self._p(doc, interlineado=1.15)
        espaciado(p, 0, 12, 1.15)
        run(p._p, texto, rpr(negrita=True, tam=12, fuente=self.fuente))

    # --- portada ---------------------------------------------------------------------
    def _portada(self, doc):
        por, inst = self.por, self.inst
        b = rpr(negrita=True, tam=12, fuente=self.fuente)
        n = rpr(tam=12, fuente=self.fuente)
        vacio = lambda: self._p(doc)  # noqa: E731

        for clave in ("universidad", "facultad", "escuela"):
            run(self._p(doc)._p, por.get(clave) or inst.get(clave, ""), b)
        vacio()
        logo = por.get("logo") or inst.get("logo_portada", "")
        if logo and (RAIZ / logo).exists():
            self._p(doc).add_run().add_picture(str(RAIZ / logo), width=Cm(4.63))
        vacio()
        run(self._p(doc)._p, por.get("tipo", "INFORME"), b)
        run(self._p(doc)._p, por.get("titulo", ""), n)
        vacio()
        etq = por.get("etiquetas", {}) or {}
        if por.get("curso"):
            run(self._p(doc)._p, etq.get("curso", "CURSO:"), b)
            run(self._p(doc)._p, por["curso"], n)
            vacio()
        docentes = [str(x) for x in a_lista(por.get("docentes")) if x]
        if docentes:
            p = self._p(doc)
            run(p._p, etq.get("docentes", "DOCENTE:"), b)
            for dname in docentes:
                br = OxmlElement("w:r")
                br.append(OxmlElement("w:br"))
                p._p.append(br)
                run(p._p, dname, n)
            vacio()
        miembros = [str(x) for x in a_lista(por.get("miembros")) if x]
        if miembros:
            run(self._p(doc)._p, etq.get("miembros", "MIEMBROS:"), b)
            for m in miembros:
                run(self._p(doc)._p, m, n)
            vacio()
        run(self._p(doc)._p, por.get("lugar", ""), b)
        run(self._p(doc)._p, por.get("anio", ""), b)
        self._salto_pagina(doc)

    # --- índice ----------------------------------------------------------------------
    def _indice(self, doc):
        ind = self.d.get("indice", {}) or {}
        if not _si(ind.get("mostrar"), True):
            return
        aviso = "Clic derecho y «Actualizar campo» para generar el índice."
        props = rpr(tam=11, fuente=self.fuente)
        self._titulo_centrado(doc, ind.get("titulo", "ÍNDICE"))
        profundidad = ind.get("profundidad", "2") or "2"
        p = doc.add_paragraph()
        campo(p._p, f'TOC \\o "1-{profundidad}" \\h \\z \\u', aviso, props)
        for clave, titulo, seq in (("figuras", "ÍNDICE DE FIGURAS", "Figura"), ("tablas", "ÍNDICE DE TABLAS", "Tabla")):
            if _si(ind.get(clave), False):
                espaciado(doc.add_paragraph(), 0, 6)
                self._titulo_centrado(doc, titulo)
                p = doc.add_paragraph()
                campo(p._p, f'TOC \\h \\z \\c "{seq}"', aviso, props)
        self._salto_pagina(doc)

    # --- pie de página ---------------------------------------------------------------
    def _pie_pagina(self, doc):
        sec = doc.sections[0]
        sec.different_first_page_header_footer = True     # la portada sin número
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        campo(p._p, "PAGE", "1", rpr(tam=10, fuente=self.fuente))

    # --- referencias -----------------------------------------------------------------
    def _referencias(self, doc):
        if not self.r.citas:
            return
        doc.add_paragraph(self.ref.get("titulo", "REFERENCIAS"), style="Heading 1")
        for k, clave in enumerate(self.r.citas):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            pf = p.paragraph_format
            pf.left_indent = Twips(567)
            pf.first_line_indent = Twips(-567)
            pf.tab_stops.add_tab_stop(Twips(567))
            espaciado(p, 0, 6, 1.15)
            run(p._p, f"[{k + 1}]\t", rpr(tam=11, fuente=self.fuente))
            for texto, cursiva in formato_ieee(self.r.bib[clave]):
                run(p._p, texto, rpr(cursiva=cursiva or None, tam=11, fuente=self.fuente))

    # --- documento -------------------------------------------------------------------
    def construir(self, destino: Path):
        doc = Document(str(PLANTILLA_DOCX))
        self._estilos(doc)

        secciones = [(s, self.r.cargar(s.get("contenido"))) for s in a_lista(self.d.get("secciones"))]
        for _, bloques in secciones:
            self.r.numerar(bloques)
        self.r.cerrar_citas()

        self._portada(doc)
        self._indice(doc)
        for k, (s, bloques) in enumerate(secciones):
            if k and _si(s.get("nueva_pagina"), False):
                self._salto_pagina(doc)
            doc.add_paragraph(s.get("titulo", ""), style="Heading 1")
            self.r.bloques(doc, bloques, 0, ANCHO)
        self._referencias(doc)
        self._pie_pagina(doc)

        doc.core_properties.title = self.por.get("titulo", "")
        doc.core_properties.author = ", ".join(str(x) for x in a_lista(self.por.get("miembros")) if x)
        try:
            doc.save(str(destino))
        except PermissionError as e:
            raise ErrorDatos(f"No se pudo guardar {destino.name}: ciérralo en Word y vuelve a intentar.") from e

        # Índice y números de figuras: los calcula Word. Si no está disponible,
        # se marca el documento para que Word los actualice al abrirlo.
        from office import actualizar_campos
        if not actualizar_campos(destino):
            doc = Document(str(destino))
            settings = doc.settings.element
            if settings.find(qn("w:updateFields")) is None:
                uf = OxmlElement("w:updateFields")
                uf.set(qn("w:val"), "true")
                settings.append(uf)
            doc.save(str(destino))


def generar_docx(datos: dict, carpeta: Path, destino: Path) -> Path:
    ConstructorAcademico(datos, carpeta).construir(destino)
    return destino
