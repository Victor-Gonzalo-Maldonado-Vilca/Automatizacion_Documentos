"""Renderizado común de contenido (typst_lite) a Word.

Cada formato de informe crea un `Renderizador` con su propio estilo (fuentes,
colores de tablas, forma de los pies de figura...) y lo usa para escribir los
bloques dentro del cuerpo del documento o dentro de una celda.
"""

from __future__ import annotations

import re
from copy import deepcopy
from pathlib import Path

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips
from docx.text.paragraph import Paragraph

import typst_lite as T
from bib import leer_bib
from datos import ErrorDatos
from resaltado import CLARO, CLARO_TEXTO, resaltar


# ============================================================================
#  Utilidades XML
# ============================================================================

def hexa(c: str) -> str:
    return c.lstrip("#").upper()


def rpr(negrita=None, cursiva=None, tam=None, color=None, fuente=None, vert=None, fondo=None) -> OxmlElement:
    """Crea un <w:rPr> en el orden que exige Word."""
    r = OxmlElement("w:rPr")
    if fuente:
        f = OxmlElement("w:rFonts")
        for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            f.set(qn(a), fuente)
        r.append(f)
    if negrita is not None:
        b = OxmlElement("w:b")
        if not negrita:
            b.set(qn("w:val"), "0")
        r.append(b)
    if cursiva is not None:
        i = OxmlElement("w:i")
        if not cursiva:
            i.set(qn("w:val"), "0")
        r.append(i)
    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), hexa(color))
        r.append(c)
    if tam:
        for tag in ("w:sz", "w:szCs"):
            s = OxmlElement(tag)
            s.set(qn("w:val"), str(int(tam * 2)))
            r.append(s)
    if fondo:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexa(fondo))
        r.append(shd)
    if vert:
        v = OxmlElement("w:vertAlign")
        v.set(qn("w:val"), "subscript" if vert == "sub" else "superscript")
        r.append(v)
    return r


def run(p_el, texto: str, props=None):
    """Agrega un run con texto (respetando tabulaciones) al párrafo XML."""
    r = OxmlElement("w:r")
    if props is not None:
        r.append(deepcopy(props))
    for k, parte in enumerate(texto.split("\t")):
        if k:
            r.append(OxmlElement("w:tab"))
        if parte:
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = parte
            r.append(t)
    p_el.append(r)
    return r


def campo(p_el, instruccion: str, resultado: str, props=None):
    """Campo de Word (PAGE, SEQ, TOC...) con un resultado provisional."""
    def fld(tipo):
        r = OxmlElement("w:r")
        if props is not None:
            r.append(deepcopy(props))
        f = OxmlElement("w:fldChar")
        f.set(qn("w:fldCharType"), tipo)
        r.append(f)
        p_el.append(r)
    fld("begin")
    r = OxmlElement("w:r")
    if props is not None:
        r.append(deepcopy(props))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instruccion} "
    r.append(it)
    p_el.append(r)
    fld("separate")
    run(p_el, resultado, props)
    fld("end")


def rpr_de(p_el):
    """rPr del primer run, o el de la marca de párrafo si no hay runs."""
    r = p_el.find(qn("w:r"))
    if r is not None and r.find(qn("w:rPr")) is not None:
        return deepcopy(r.find(qn("w:rPr")))
    ppr = p_el.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:rPr")) is not None:
        return deepcopy(ppr.find(qn("w:rPr")))
    return None


def vaciar(p_el):
    for hijo in list(p_el):
        if hijo.tag != qn("w:pPr"):
            p_el.remove(hijo)


def poner_texto(p_el, texto: str, props=None):
    """Reemplaza el texto de un párrafo conservando su formato."""
    props = props if props is not None else rpr_de(p_el)
    vaciar(p_el)
    run(p_el, texto, props)


def sombrear(tc, color: str):
    tcpr = tc.get_or_add_tcPr()
    for v in tcpr.findall(qn("w:shd")):
        tcpr.remove(v)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexa(color))
    tcpr.append(shd)


def bordes_tabla(tbl, color="000000", tam=4):
    tblpr = tbl.tblPr
    for v in tblpr.findall(qn("w:tblBorders")):
        tblpr.remove(v)
    bordes = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{lado}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), str(tam))
        b.set(qn("w:space"), "0")
        b.set(qn("w:color"), hexa(color))
        bordes.append(b)
    tblpr.append(bordes)


def tabla_props(tbl, ancho: int, sangria: int, margen_v: int = 40, margen_h: int = 108):
    tblpr = tbl.tblPr
    for tag in ("w:tblW", "w:tblInd", "w:tblCellMar", "w:tblLayout", "w:jc"):
        for v in tblpr.findall(qn(tag)):
            tblpr.remove(v)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), str(ancho))
    w.set(qn("w:type"), "dxa")
    ind = OxmlElement("w:tblInd")
    ind.set(qn("w:w"), str(sangria))
    ind.set(qn("w:type"), "dxa")
    lay = OxmlElement("w:tblLayout")
    lay.set(qn("w:type"), "fixed")
    mar = OxmlElement("w:tblCellMar")
    for lado, v in (("top", margen_v), ("left", margen_h), ("bottom", margen_v), ("right", margen_h)):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    for el in (w, ind, lay, mar):
        tblpr.append(el)


def ancho_celda(celda, ancho: int):
    tcpr = celda._tc.get_or_add_tcPr()
    for v in tcpr.findall(qn("w:tcW")):
        tcpr.remove(v)
    w = OxmlElement("w:tcW")
    w.set(qn("w:w"), str(ancho))
    w.set(qn("w:type"), "dxa")
    tcpr.insert(0, w)


def espaciado(p: Paragraph, antes=0, despues=0, interlineado=None):
    pf = p.paragraph_format
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)
    if interlineado is not None:
        pf.line_spacing = interlineado


def borde_imagen(r):
    for sppr in r._r.iter(qn("pic:spPr")):
        ln = OxmlElement("a:ln")
        ln.set("w", "6350")
        fill = OxmlElement("a:solidFill")
        clr = OxmlElement("a:srgbClr")
        clr.set("val", "000000")
        fill.append(clr)
        ln.append(fill)
        sppr.append(ln)


# ============================================================================
#  Renderizador
# ============================================================================

ESTILO_BASE = {
    "fuente": None,              # None = la del documento base
    "fuente_codigo": "Consolas",
    "tam": 11,
    "interlineado": 1.08,
    "despues": 8,                # espacio después de cada párrafo (pt)
    "sangria_primera": 0,        # sangría de primera línea de párrafos (twips)
    "color_enc": "#C8310E",      # encabezado de tablas y título de código
    "color_texto_enc": "#FFFFFF",
    "borde_tabla": ("#000000", 4),
    "tam_tabla": 11,
    "alinear_enc": "centro",     # alineación del encabezado de tablas: "centro" | "izquierda"
    "margen_tabla": (40, 108),   # (vertical, horizontal) en twips
    "tam_pie": 10,
    "pie": "punto",              # "punto": "Figura 1. Texto" | "apa": "**Figura 1** _Texto_"
    "tabla_pie": "centro",       # "centro": centrado | "izquierda_negrita": "Tabla 1. Texto" en negrita
    "nota_borde": "#C8310E",
    "nota_fondo": "#FBEFEC",
    "estilos_titulo": None,      # {nivel: styleId} para usar estilos de título de Word
    "campos_seq": False,         # numerar pies con campos SEQ (para índices de figuras/tablas)
    "vinetas": ("•", "–"),
}

NOMBRES = {"figura": "Figura", "tabla": "Tabla", "codigo": "Código", "ecuacion": "Ecuación"}


class Renderizador:
    def __init__(self, carpeta: Path, estilo: dict, archivo_bib: str = "", incluir_todas: bool = False):
        self.carpeta = carpeta
        self.e = {**ESTILO_BASE, **estilo}
        self.contadores = {k: 0 for k in NOMBRES}
        self.etiquetas: dict[str, str] = {}
        self.citas: list[str] = []
        self.bib: dict = {}
        self.incluir_todas = incluir_todas
        if archivo_bib:
            ruta = carpeta / archivo_bib
            if ruta.exists():
                self.bib = {e["clave"]: e for e in leer_bib(ruta)}

    # --- carga y numeración ----------------------------------------------------
    def leer(self, ruta: str) -> str:
        p = self.carpeta / ruta
        if not p.exists():
            raise ErrorDatos(f"No existe el archivo: {p}")
        return p.read_text(encoding="utf-8")

    def cargar(self, valor) -> list[dict]:
        if not valor:
            return []
        v = str(valor).strip()
        texto = self.leer(v) if v.endswith(".typ") else v
        return T.parsear(texto, self.leer)

    def numerar(self, bloques: list[dict]):
        """Primera pasada: numera figuras, tablas, código y ecuaciones; ordena citas."""
        for b in bloques:
            tipo = b["t"]
            if tipo in NOMBRES and (tipo == "ecuacion" or b.get("pie")):
                self.contadores[tipo] += 1
                b["numero"] = self.contadores[tipo]
                if b.get("etiqueta"):
                    self.etiquetas[b["etiqueta"]] = f"{NOMBRES[tipo]} {b['numero']}"
            inlines = []
            if tipo in ("titulo", "parrafo", "nota"):
                inlines.append(b["inl"])
            if tipo == "lista":
                inlines += [it[2] for it in b["items"]]
            if tipo == "tabla":
                inlines += [b.get("pie")] + b["encabezado"] + b["celdas"]
            if tipo in ("figura", "codigo"):
                inlines.append(b.get("pie"))
            for inl in inlines:
                for x in inl or []:
                    c = x.get("cita")
                    if c and c in self.bib and c not in self.citas:
                        self.citas.append(c)

    def cerrar_citas(self):
        if self.incluir_todas:
            self.citas += [k for k in self.bib if k not in self.citas]

    # --- inline -----------------------------------------------------------------
    def props(self, **kw):
        kw.setdefault("fuente", self.e["fuente"])
        return rpr(**kw)

    def inline(self, p_el, inl: list[dict], tam=None, color=None, cursiva_base=False, negrita_base=False):
        tam = tam or self.e["tam"]
        k = 0
        while k < len(inl):
            x = inl[k]
            if "salto" in x:
                r = OxmlElement("w:r")
                r.append(OxmlElement("w:br"))
                p_el.append(r)
            elif "math" in x:
                for texto, vert, it in T.math_segmentos(x["math"]):
                    run(p_el, texto, rpr(cursiva=it, tam=tam, fuente="Cambria Math", vert=vert or None, color=color))
            elif "cita" in x:
                clave = x["cita"]
                if clave in self.etiquetas:
                    texto = self.etiquetas[clave]
                elif clave in self.bib:
                    texto = f"[{self.citas.index(clave) + 1}]"
                    entre = inl[k + 1] if k + 1 < len(inl) else None
                    sig = inl[k + 2] if k + 2 < len(inl) else None
                    if sig and sig.get("cita") in self.bib and entre and entre.get("txt", "x").strip() == "":
                        texto += ", "  # [2] [3] -> [2], [3]
                        k += 1
                else:
                    texto = f"[{clave}?]"
                run(p_el, texto, self.props(tam=tam, color=color, cursiva=cursiva_base or None,
                                            negrita=negrita_base or None))
            elif x.get("code"):
                run(p_el, x["txt"], rpr(tam=tam - 1, fuente=self.e["fuente_codigo"], fondo="#F0F0F0", color=color))
            else:
                run(p_el, x["txt"], self.props(negrita=(x["b"] or negrita_base) or None,
                                               cursiva=(x["i"] != cursiva_base) or None, tam=tam, color=color))
            k += 1

    # --- párrafos y pies ---------------------------------------------------------
    def parrafo(self, cont, sangria: int, alinear=WD_ALIGN_PARAGRAPH.JUSTIFY) -> Paragraph:
        p = cont.add_paragraph()
        p.paragraph_format.left_indent = Twips(sangria)
        p.alignment = alinear
        espaciado(p, 0, self.e["despues"], self.e["interlineado"])
        return p

    def pie(self, cont, sangria, tipo, numero, inl, despues=None):
        e = self.e
        despues = e["despues"] if despues is None else despues
        izquierda = tipo == "tabla" and e["tabla_pie"] == "izquierda_negrita"
        p = self.parrafo(cont, sangria, WD_ALIGN_PARAGRAPH.LEFT if izquierda else WD_ALIGN_PARAGRAPH.CENTER)
        espaciado(p, 3, despues, 1.0)
        tam = e["tam_pie"]
        negrita = self.props(negrita=True, tam=tam)
        nombre = NOMBRES[tipo]
        run(p._p, f"{nombre} ", negrita)
        if e["campos_seq"]:
            campo(p._p, f"SEQ {nombre} \\* ARABIC", str(numero), negrita)
        else:
            run(p._p, str(numero), negrita)
        if izquierda:
            run(p._p, ". ", negrita)
            self.inline(p._p, inl, tam=tam, negrita_base=True)
        elif e["pie"] == "apa":
            run(p._p, " ", negrita)
            self.inline(p._p, inl, tam=tam, cursiva_base=True)
        else:
            run(p._p, ". ", negrita)
            self.inline(p._p, inl, tam=tam)
        return p

    def separador(self, cont):
        p = cont.add_paragraph()
        espaciado(p, 0, 6, 1.0)
        p.paragraph_format.line_spacing = Pt(4)

    # --- bloques -------------------------------------------------------------------
    def bloques(self, cont, bloques: list[dict], sangria: int, ancho: int):
        for b in bloques:
            getattr(self, "b_" + b["t"])(cont, b, sangria, ancho)

    def b_titulo(self, cont, b, sangria, ancho):
        estilos = self.e["estilos_titulo"]
        if estilos:
            p = cont.add_paragraph(style=estilos.get(b["nivel"], estilos[max(estilos)]))
            self.inline(p._p, b["inl"])
            return
        p = self.parrafo(cont, sangria, WD_ALIGN_PARAGRAPH.LEFT)
        espaciado(p, 6, 3, self.e["interlineado"])
        p.paragraph_format.keep_with_next = True
        self.inline(p._p, b["inl"], negrita_base=True, cursiva_base=b["nivel"] >= 2)

    def b_parrafo(self, cont, b, sangria, ancho):
        p = self.parrafo(cont, sangria)
        if self.e["sangria_primera"]:
            p.paragraph_format.first_line_indent = Twips(self.e["sangria_primera"])
        self.inline(p._p, b["inl"])

    def b_lista(self, cont, b, sangria, ancho):
        contadores: dict[int, int] = {}
        vinetas = self.e["vinetas"]
        for k, (nivel, ordenada, inl) in enumerate(b["items"]):
            for n in [n for n in contadores if n > nivel]:
                del contadores[n]
            contadores[nivel] = contadores.get(nivel, 0) + 1
            p = self.parrafo(cont, sangria + 360 * (nivel + 1))
            p.paragraph_format.first_line_indent = Twips(-300)
            espaciado(p, 0, self.e["despues"] if k == len(b["items"]) - 1 else 2, self.e["interlineado"])
            marca = f"{contadores[nivel]}.\t" if ordenada else f"{vinetas[min(nivel, len(vinetas) - 1)]}\t"
            run(p._p, marca, self.props(tam=self.e["tam"]))
            self.inline(p._p, inl)

    def b_figura(self, cont, b, sangria, ancho):
        ruta = self.carpeta / (b["ruta"] or "")
        if not b["ruta"] or not ruta.exists():
            raise ErrorDatos(f"No existe la imagen: {ruta}")
        p = self.parrafo(cont, sangria, WD_ALIGN_PARAGRAPH.CENTER)
        espaciado(p, 6, 0 if b.get("pie") else self.e["despues"], 1.0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run()
        r.add_picture(str(ruta), width=Twips(int(ancho * b["ancho"])))
        if b.get("borde"):
            borde_imagen(r)
        if b.get("pie"):
            self.pie(cont, sangria, "figura", b["numero"], b["pie"])

    def _anchos(self, columnas: list[str], celdas: list, encabezado: list, ancho: int) -> list[int]:
        n = len(columnas)
        filas = ([encabezado] if encabezado else []) + [celdas[k:k + n] for k in range(0, len(celdas), n)]
        largo = lambda inl: sum(len(x.get("txt", "")) for x in inl)  # noqa: E731
        margen_h = self.e["margen_tabla"][1]
        fijos, fr = {}, {}
        for c, spec in enumerate(columnas):
            m = re.match(r"([\d.]+)\s*(cm|mm|pt|in|fr|%)?", spec)
            if spec == "auto" or not m:
                maximo = max((largo(f[c]) for f in filas if c < len(f)), default=4)
                fijos[c] = min(int(maximo * 115 + 2 * margen_h + 60), int(ancho * 0.45))
            elif m.group(2) == "fr":
                fr[c] = float(m.group(1))
            elif m.group(2) == "%":
                fijos[c] = int(ancho * float(m.group(1)) / 100)
            else:
                factor = {"cm": 567, "mm": 56.7, "pt": 20, "in": 1440}.get(m.group(2) or "pt", 20)
                fijos[c] = int(float(m.group(1)) * factor)
        resto = max(ancho - sum(fijos.values()), 600 * max(len(fr), 1))
        total_fr = sum(fr.values()) or 1
        anchos = [fijos.get(c, int(resto * fr.get(c, 0) / total_fr)) for c in range(n)]
        if not fr:  # todo auto: escalar al ancho disponible
            escala = ancho / max(sum(anchos), 1)
            anchos = [int(a * escala) for a in anchos]
        return anchos

    def b_tabla(self, cont, b, sangria, ancho):
        e = self.e
        if b.get("pie"):
            p = self.pie(cont, sangria, "tabla", b["numero"], b["pie"], despues=3)
            p.paragraph_format.keep_with_next = True
        n = len(b["columnas"])
        filas = [b["celdas"][k:k + n] for k in range(0, len(b["celdas"]), n)]
        if b["encabezado"]:
            filas.insert(0, b["encabezado"])
        tabla = cont.add_table(rows=len(filas), cols=n)
        tbl = tabla._tbl
        tabla_props(tbl, ancho, sangria, *e["margen_tabla"])
        bordes_tabla(tbl, *e["borde_tabla"])
        anchos = self._anchos(b["columnas"], b["celdas"], b["encabezado"], ancho)
        for k, gc in enumerate(tbl.tblGrid.findall(qn("w:gridCol"))):
            gc.set(qn("w:w"), str(anchos[k]))
        for r, fila in enumerate(filas):
            es_enc = bool(b["encabezado"]) and r == 0
            if es_enc:
                tabla.rows[r]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
            for c in range(n):
                cel = tabla.cell(r, c)
                ancho_celda(cel, anchos[c])
                cel.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p = cel.paragraphs[0]
                espaciado(p, 0, 0, 1.0)
                inl = fila[c] if c < len(fila) else []
                if es_enc:
                    sombrear(cel._tc, e["color_enc"])
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if e["alinear_enc"] == "centro" else WD_ALIGN_PARAGRAPH.LEFT
                    self.inline(p._p, inl, tam=e["tam_tabla"], color=e["color_texto_enc"], negrita_base=True)
                else:
                    self.inline(p._p, inl, tam=e["tam_tabla"])
        self.separador(cont)

    def b_codigo(self, cont, b, sangria, ancho):
        e = self.e
        tiene_titulo = bool(b.get("titulo"))
        tabla = cont.add_table(rows=2 if tiene_titulo else 1, cols=1)
        tbl = tabla._tbl
        tabla_props(tbl, ancho, sangria, margen_v=60)
        bordes_tabla(tbl, "A0A0A0")
        tbl.tblGrid.find(qn("w:gridCol")).set(qn("w:w"), str(ancho))
        fila = 0
        if tiene_titulo:
            cel = tabla.cell(0, 0)
            ancho_celda(cel, ancho)
            sombrear(cel._tc, e["color_enc"])
            p = cel.paragraphs[0]
            espaciado(p, 0, 0, 1.0)
            run(p._p, b["titulo"], rpr(negrita=True, tam=9, color=e["color_texto_enc"], fuente=e["fuente_codigo"]))
            fila = 1
        cel = tabla.cell(fila, 0)
        ancho_celda(cel, ancho)
        sombrear(cel._tc, "#F7F7F7")
        lineas = resaltar(b["codigo"], b.get("lenguaje", ""), b.get("titulo") or "", CLARO, CLARO_TEXTO) or [[]]
        ancho_num = len(str(len(lineas)))
        for k, segmentos in enumerate(lineas):
            p = cel.paragraphs[0] if k == 0 else cel.add_paragraph()
            espaciado(p, 0, 0, 1.0)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if len(lineas) > 1:
                run(p._p, f"{k + 1:>{ancho_num}}  ", rpr(tam=9, color="#8C8C8C", fuente=e["fuente_codigo"]))
            for texto, color, cursiva in segmentos:
                run(p._p, texto.expandtabs(4), rpr(tam=9, color=color, cursiva=cursiva or None, fuente=e["fuente_codigo"]))
        if b.get("pie"):
            self.pie(cont, sangria, "codigo", b["numero"], b["pie"])
        else:
            self.separador(cont)

    def b_nota(self, cont, b, sangria, ancho):
        e = self.e
        p = self.parrafo(cont, sangria + 60)
        ppr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        for lado, color, tam in (("top", "BEBEBE", 4), ("left", hexa(e["nota_borde"]), 18),
                                 ("bottom", "BEBEBE", 4), ("right", "BEBEBE", 4)):
            el = OxmlElement(f"w:{lado}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(tam))
            el.set(qn("w:space"), "4")
            el.set(qn("w:color"), color)
            bdr.append(el)
        ppr.append(bdr)
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexa(e["nota_fondo"]))
        ppr.append(shd)
        p.paragraph_format.right_indent = Twips(80)
        espaciado(p, 6, 10, e["interlineado"])
        self.inline(p._p, b["titulo"], color=e["nota_borde"], negrita_base=True)
        run(p._p, ": ", self.props(negrita=True, color=e["nota_borde"]))
        self.inline(p._p, b["inl"])

    def b_ecuacion(self, cont, b, sangria, ancho):
        p = self.parrafo(cont, sangria, WD_ALIGN_PARAGRAPH.LEFT)
        espaciado(p, 4, self.e["despues"], self.e["interlineado"])
        tabs = p.paragraph_format.tab_stops
        tabs.add_tab_stop(Twips(sangria + ancho // 2), WD_TAB_ALIGNMENT.CENTER)
        tabs.add_tab_stop(Twips(sangria + ancho), WD_TAB_ALIGNMENT.RIGHT)
        run(p._p, "\t")
        for texto, vert, it in T.math_segmentos(b["math"]):
            run(p._p, texto, rpr(cursiva=it, tam=self.e["tam"] + 1, fuente="Cambria Math", vert=vert or None))
        run(p._p, f"\t({b['numero']})", self.props(tam=self.e["tam"]))
